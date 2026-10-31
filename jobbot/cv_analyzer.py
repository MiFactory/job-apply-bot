"""Read a CV (PDF/DOCX), extract facts and propose job-search criteria."""
import re, pathlib, collections

SKILLS = {
 "sales": ["sales", "vente", "business development", "account manager", "key account", "b2b", "crm", "salesforce", "hubspot", "négociation", "negotiation", "prospection", "lead generation", "tender", "pre-sales", "presales"],
 "it": ["it support", "helpdesk", "service desk", "windows", "linux", "active directory", "azure", "microsoft 365", "intune", "network", "réseau", "cybersecurity", "cybersécurité", "sysadmin", "devops", "cloud", "itil", "sql", "python", "javascript", "java", "c#", "docker", "kubernetes", "vmware", "firewall"],
 "project": ["project manager", "chef de projet", "scrum", "agile", "pmo", "prince2", "pmp", "kanban", "jira", "product owner", "stakeholder", "budget", "planning"],
 "data": ["data analyst", "power bi", "tableau", "etl", "machine learning", "bi ", "analytics"],
 "embedded": ["embedded", "firmware", "plc", "automation", "automatisme", "iot", "scada", "opc"],
 "retail": ["retail", "store manager", "customer service", "service client", "merchandising"],
}
ROLE_FAMILIES = {
 "sales": ["Business Developer", "Account Manager", "Sales Executive", "Key Account Manager", "Sales Engineer", "Inside Sales", "Commercial IT", "Presales"],
 "it": ["IT Support", "IT Consultant", "System Administrator", "Service Desk", "IT Project Manager", "Solutions Architect", "Cloud Engineer", "Cybersecurity"],
 "project": ["Project Manager", "Chef de projet", "Product Owner", "Scrum Master", "Programme Manager"],
 "data": ["Data Analyst", "BI Analyst", "Data Engineer"],
 "embedded": ["Embedded Engineer", "Automation Engineer", "Software Engineer"],
 "retail": ["Customer Service", "Retail Manager", "Network Manager"],
}
LANGS = {"english": ["english", "anglais"], "french": ["french", "français", "francais"], "dutch": ["dutch", "néerlandais", "neerlandais", "nederlands"],
         "spanish": ["spanish", "espagnol"], "german": ["german", "allemand"], "arabic": ["arabic", "arabe"]}


def read_text(path) -> str:
    p = pathlib.Path(path)
    if p.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        return "\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages)
    if p.suffix.lower() == ".docx":
        import docx
        d = docx.Document(str(p))
        parts = [x.text for x in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts += [c.text for c in row.cells]
        return "\n".join(parts)
    return p.read_text(encoding="utf-8", errors="ignore")


def analyze(path) -> dict:
    txt = read_text(path)
    low = txt.lower()
    email = (re.findall(r"[\w.+-]+@[\w-]+\.[\w.]+", txt) or [""])[0]
    phone = (re.findall(r"(\+?\d[\d \-.()]{8,16}\d)", txt) or [""])[0].strip()
    lines = [l.strip() for l in txt.splitlines() if l.strip()]
    name = lines[0] if lines else ""
    fam_score = collections.Counter()
    found = []
    for fam, kws in SKILLS.items():
        for k in kws:
            n = low.count(k)
            if n:
                fam_score[fam] += n
                found.append(k.strip())
    years = [int(y) for y in re.findall(r"(\d{1,2})\s*\+?\s*(?:years|ans|jaar)", low)]
    langs = [l for l, ks in LANGS.items() if any(k in low for k in ks)]
    top = [f for f, _ in fam_score.most_common(3)]
    titles = []
    for f in top:
        titles += ROLE_FAMILIES.get(f, [])
    return {"name": name, "email": email, "phone": phone, "years_experience": (sorted(years)[len(years)//2] if years else None),
            "languages": langs, "skills": sorted(set(found)), "families": top, "suggested_titles": titles, "chars": len(txt)}
